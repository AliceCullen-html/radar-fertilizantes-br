"""Fase 6: fontes externas para comparar o Brasil com o mercado internacional.

- Banco Mundial, Commodity Price Data (The Pink Sheet), arquivo mensal
  CMO-Historical-Data-Monthly.xlsx, licença CC BY 4.0: ureia, DAP, TSP, cloreto de potássio e
  rocha fosfática, em US$/t nominais. A URL do arquivo muda a cada edição, então o programa a
  procura na página Commodity Markets. O arquivo vem sempre inteiro (desde 1960) e o Banco
  Mundial revisa meses passados: cada ingestão regrava a série desde 2015-01 e lista o que
  foi revisado. Grava data/raw_benchmark/pink_sheet_mensal.csv (um mês por linha) e
  data/raw_benchmark/fonte.json (URL, data de atualização e sha256 do arquivo lido).
- Banco Central do Brasil, SGS 3698 (taxa de câmbio livre, dólar americano, venda, média do
  período, mensal; é a média das PTAX de venda dos dias úteis), dados abertos sob ODbL.
  Incremental, como o Comex Stat: só consulta os meses que faltam, até o último publicado.
  Grava data/raw_cambio/ptax_venda_media_mensal.csv.

As duas fontes usam o mesmo cliente: pausa entre chamadas, novas tentativas com espera
exponencial em HTTP 429, 5xx, falha de rede (inclusive corte no meio do corpo e erro TLS
transitório) e corpo inesperado (o firewall do BCB às vezes responde com uma página HTML e
HTTP 200), timeout, limite de tamanho do corpo e verificação TLS sempre ligada. Nenhuma das
duas pede chave.

Uso (pela CLI da ingestão):
    python -m src.ingest --inicio 2015-01 --serie benchmark
    python -m src.ingest --inicio 2015-01 --serie cambio

Fonte: Banco Mundial (Pink Sheet); BCB/PTAX (SGS 3698).
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import random
import re
import ssl
import time
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlsplit

import pandas as pd
import requests
from urllib3.util import parse_url

from . import config

log = logging.getLogger(__name__)

COLUNAS_PINK = ["mes", *config.PINK_SERIES.values()]
COLUNAS_PTAX = ["mes", "brl_por_usd"]
_RE_PERIODO_PINK = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
_RE_ATUALIZADO = re.compile(r"Updated on\s+([A-Za-z]+ \d{1,2}, \d{4})")
_RE_DATA_SGS = re.compile(r"^01/(0[1-9]|1[0-2])/(\d{4})$")


class ErroFonte(RuntimeError):
    """Falha ao obter ou ler o Banco Mundial ou o BCB."""


class _CorpoInesperado(Exception):
    """Resposta 200 com corpo que não é o esperado (ex.: página HTML do firewall)."""


# --- Cliente HTTP ---------------------------------------------------------------------


class ClienteFontes:
    """GET nas fontes externas, com o mesmo padrão do cliente do Comex Stat.

    - Pausa mínima entre chamadas; novas tentativas com espera exponencial em HTTP 429, 5xx,
      falha de rede e corpo inesperado, usando Retry-After quando houver.
    - Verificação TLS sempre ligada (os três hosts usam CAs públicas que o certifi valida).
    - Só acessa os hosts de config.HOSTS_FONTES, por https, inclusive em redirecionamentos.
    """

    def __init__(
        self,
        *,
        sessao: requests.Session | None = None,
        pausa: float = config.PAUSA_FONTES_S,
        max_tentativas: int = config.MAX_TENTATIVAS,
        dormir=time.sleep,
        relogio=time.monotonic,
    ):
        self.sessao = sessao if sessao is not None else requests.Session()
        self.sessao.headers.update({"User-Agent": config.USER_AGENT})
        self.pausa = pausa
        self.max_tentativas = max_tentativas
        self._dormir = dormir
        self._relogio = relogio
        self._ultima_chamada: float | None = None

    def get(self, url: str, *, esperado: str, params: dict | None = None) -> requests.Response:
        """`esperado`: "html", "xlsx" ou "json" (lista). Devolve a resposta já conferida."""
        motivo = ""
        for tentativa in range(1, self.max_tentativas + 1):
            espera = self._backoff(tentativa)
            try:
                resp = self._seguir(url, params, esperado)
                if resp.status_code == 200:
                    _ler_corpo(resp, config.PINK_MAX_BYTES if esperado == "xlsx" else config.MAX_BYTES_PAGINA)
                    _conferir_corpo(resp, esperado)
                    return resp
                if resp.status_code != 429 and resp.status_code < 500:
                    raise ErroFonte(f"HTTP {resp.status_code} em {_sem_query(url)}")
                motivo = f"HTTP {resp.status_code}"
                sugerida = resp.headers.get("Retry-After", "")
                if sugerida.isdigit():  # com teto: um Retry-After de uma hora não segura o job até o timeout
                    espera = max(espera, min(float(sugerida), config.BACKOFF_MAX_S))
            except _CorpoInesperado as exc:
                motivo = str(exc)
            except requests.exceptions.SSLError as exc:
                if _falha_de_certificado(exc):
                    raise ErroFonte(f"Falha na verificação TLS de {urlsplit(url).hostname}: {exc}") from exc
                motivo = f"erro TLS transitório ({exc})"
            except (requests.ConnectionError, requests.Timeout, requests.exceptions.ChunkedEncodingError, requests.exceptions.ContentDecodingError) as exc:
                motivo = f"falha de rede ({exc})"
            if tentativa < self.max_tentativas:
                log.warning("%s em %s; tentativa %d/%d em %.0f s", motivo, _sem_query(url), tentativa + 1, self.max_tentativas, espera)
                self._dormir(espera)
        raise ErroFonte(f"GET {_sem_query(url)} falhou após {self.max_tentativas} tentativas: {motivo}")

    def _seguir(self, url: str, params: dict | None, esperado: str) -> requests.Response:
        for _ in range(4):  # até 3 redirecionamentos, sempre dentro dos hosts permitidos
            _conferir_host(url, params)
            self._aguardar_pausa()
            resp = self.sessao.request(
                "GET", url, params=params, headers={"Accept": _ACCEPT[esperado]}, timeout=config.TIMEOUT_S,
                allow_redirects=False, verify=True, stream=True,
            )
            if resp.status_code not in (301, 302, 303, 307, 308):
                return resp
            try:
                url, params = urljoin(url, resp.headers.get("Location", "")), None
            except ValueError as exc:
                raise ErroFonte(f"redirecionamento inválido a partir de {_sem_query(url)}") from exc
        raise ErroFonte(f"redirecionamentos demais a partir de {_sem_query(url)}")

    def _aguardar_pausa(self) -> None:
        if self._ultima_chamada is not None:
            falta = self.pausa - (self._relogio() - self._ultima_chamada)
            if falta > 0:
                self._dormir(falta)
        self._ultima_chamada = self._relogio()

    @staticmethod
    def _backoff(tentativa: int) -> float:
        return min(config.BACKOFF_MAX_S, config.BACKOFF_BASE_S * 2 ** (tentativa - 1)) + random.uniform(0, 1)


# O firewall do BCB recusa (com uma página HTML) pedidos sem Accept de JSON.
_ACCEPT = {
    "json": "application/json",
    "html": "text/html,application/xhtml+xml",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*;q=0.5",
}


def _conferir_host(url: str, params: dict | None = None) -> None:
    """Só https e só os hosts permitidos, conferidos na URL que o requests vai de fato usar:
    "https://outro.example\\@api.bcb.gov.br/" parece do BCB para o urlsplit, mas o requests a
    envia para outro.example. Barra invertida, "@", espaço e caractere de controle reprovam."""
    try:
        netloc = urlsplit(url).netloc
    except ValueError as exc:
        raise ErroFonte(f"endereço inválido: {url[:200]!r}") from exc
    if re.search(r"[\\@\s\x00-\x1f\x7f]", netloc) or "\\" in url:
        raise ErroFonte(f"endereço com caractere não permitido: {_sem_query(url)!r}")
    try:
        preparada = requests.Request("GET", url, params=params).prepare().url
        hosts = {urlsplit(url).hostname, urlsplit(preparada).hostname, parse_url(preparada).host}
        esquemas = {urlsplit(url).scheme, urlsplit(preparada).scheme}
    except (ValueError, requests.exceptions.RequestException) as exc:
        raise ErroFonte(f"endereço inválido: {_sem_query(url)!r}") from exc
    if esquemas != {"https"} or len(hosts) != 1 or not hosts <= set(config.HOSTS_FONTES):
        raise ErroFonte(f"endereço fora das fontes permitidas: {_sem_query(url)}")


def _falha_de_certificado(exc: BaseException) -> bool:
    """Erro de verificação do certificado (aborta) ou erro TLS transitório (repete)."""
    vistos = set()
    while exc is not None and id(exc) not in vistos:
        vistos.add(id(exc))
        if isinstance(exc, ssl.SSLCertVerificationError) or "CERTIFICATE_VERIFY_FAILED" in str(exc):
            return True
        exc = exc.__cause__ or exc.__context__ or (exc.args[0] if exc.args and isinstance(exc.args[0], BaseException) else None)
    return False


def _ler_corpo(resp, limite: int) -> None:
    """Lê o corpo em blocos e para no limite: um arquivo enorme não chega inteiro à memória."""
    tamanho = str(resp.headers.get("Content-Length", ""))
    if tamanho.isdigit() and int(tamanho) > limite:
        raise ErroFonte(f"a resposta tem {tamanho} bytes, acima do limite de {limite}")
    if hasattr(resp, "iter_content") and not getattr(resp, "_content_consumed", False):
        partes, total = [], 0
        for parte in resp.iter_content(chunk_size=1 << 16):
            total += len(parte)
            if total > limite:
                resp.close()
                raise ErroFonte(f"a resposta passou de {limite} bytes; nada foi lido além disso")
            partes.append(parte)
        resp._content, resp._content_consumed = b"".join(partes), True
    elif len(resp.content) > limite:
        raise ErroFonte(f"a resposta tem {len(resp.content)} bytes, acima do limite de {limite}")


def _conferir_corpo(resp: requests.Response, esperado: str) -> None:
    corpo = resp.content
    if esperado == "xlsx":
        if not corpo.startswith(b"PK\x03\x04"):  # xlsx é um zip
            raise _CorpoInesperado("o arquivo da Pink Sheet não veio como xlsx")
    elif esperado == "json":
        try:
            dados = json.loads(corpo.decode("utf-8"))
        except (UnicodeDecodeError, ValueError):
            pagina = corpo.lstrip()[:200].lower()
            if pagina.startswith((b"<", b"<?xml")):
                raise _CorpoInesperado("o BCB devolveu uma página HTML no lugar do JSON (firewall)") from None
            raise _CorpoInesperado("resposta do BCB que não é JSON") from None
        if not isinstance(dados, list):
            raise _CorpoInesperado("resposta do BCB que não é uma lista")
    elif esperado == "html":
        if b"<html" not in corpo[:5000].lower():
            raise _CorpoInesperado("a página do Banco Mundial não veio como HTML")


def _sem_query(url: str) -> str:
    try:
        p = urlsplit(url)
    except ValueError:
        return url[:200]
    return f"{p.scheme}://{p.netloc}{p.path}"


def _gravar_csv(df: pd.DataFrame, caminho: Path) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    temporario = caminho.with_name(caminho.name + ".tmp")
    df.to_csv(temporario, index=False, encoding="utf-8", lineterminator="\n")
    os.replace(temporario, caminho)


def _gravar_texto(caminho: Path, texto: str) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    temporario = caminho.with_name(caminho.name + ".tmp")
    temporario.write_text(texto, encoding="utf-8")
    os.replace(temporario, caminho)


def _meses(inicio: str, fim: str) -> list[str]:
    ano, mes = map(int, inicio.split("-"))
    ano_fim, mes_fim = map(int, fim.split("-"))
    meses = []
    while (ano, mes) <= (ano_fim, mes_fim):
        meses.append(f"{ano:04d}-{mes:02d}")
        ano, mes = (ano + 1, 1) if mes == 12 else (ano, mes + 1)
    return meses


# --- Banco Mundial: Pink Sheet --------------------------------------------------------


def url_do_arquivo_pink(html: str, pagina: str = config.PINK_PAGINA) -> str:
    """O link do xlsx mensal na página Commodity Markets: um só, no host do Banco Mundial."""
    hrefs = re.findall(r"""href\s*=\s*["']([^"']+)["']""", html, re.IGNORECASE)
    candidatos = sorted(
        {
            urljoin(pagina, h)
            for h in hrefs
            if urlsplit(urljoin(pagina, h)).path.endswith("/" + config.PINK_NOME_ARQUIVO)
        }
    )
    validos = [u for u in candidatos if _host_seguro(u) == config.PINK_HOST_ARQUIVO]
    if len(validos) != 1:
        raise ErroFonte(
            f"a página {config.PINK_PAGINA} deveria ter um link para {config.PINK_NOME_ARQUIVO} em"
            f" {config.PINK_HOST_ARQUIVO}, mas tem {len(validos)} (links com esse nome: {len(candidatos)})."
            " O layout da página pode ter mudado."
        )
    return validos[0]


def _host_seguro(url: str) -> str | None:
    """O host da URL se ela passar em _conferir_host; senão None."""
    try:
        _conferir_host(url)
    except ErroFonte:
        return None
    return urlsplit(url).hostname


def _nome_coluna(valor) -> str:
    return str(valor or "").strip().strip("*").strip()


def ler_xlsx_pink(conteudo: bytes) -> tuple[pd.DataFrame, str | None]:
    """(um mês por linha, com as cinco séries em US$/t; data "Updated on" em AAAA-MM-DD).

    Confere o layout: a aba, o nome e a unidade ($/mt) de cada coluna e o formato AAAAMmm dos
    meses. Valores ausentes ("..", "…" ou vazio) viram NaN. Qualquer surpresa é ErroFonte.
    """
    from openpyxl import load_workbook  # só a ingestão do Banco Mundial precisa

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
            descompactado = sum(info.file_size for info in arquivo.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise ErroFonte(f"não foi possível abrir o xlsx da Pink Sheet ({type(exc).__name__})") from exc
    if descompactado > config.PINK_MAX_DESCOMPACTADO:
        raise ErroFonte(f"o xlsx da Pink Sheet descompacta para {descompactado} bytes, acima do limite de {config.PINK_MAX_DESCOMPACTADO}")
    try:
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:  # zip ou XML inválido: o arquivo não é o esperado
        raise ErroFonte(f"não foi possível abrir o xlsx da Pink Sheet ({type(exc).__name__})") from exc
    try:
        if config.PINK_ABA not in livro.sheetnames:
            raise ErroFonte(f"a Pink Sheet não tem a aba {config.PINK_ABA!r} (abas: {livro.sheetnames})")
        folha = livro[config.PINK_ABA]
        folha.reset_dimensions()  # não confia na dimensão gravada na aba, que pode cortar linhas
        linhas = list(folha.iter_rows(values_only=True))
    finally:
        livro.close()

    atualizado = None
    for linha in linhas[:10]:
        achado = next((_RE_ATUALIZADO.search(str(c)) for c in linha if c and _RE_ATUALIZADO.search(str(c))), None)
        if achado:
            try:
                atualizado = datetime.strptime(achado.group(1), "%B %d, %Y").strftime("%Y-%m-%d")
            except ValueError:  # "Sept 02, 2026": a data é só registro, não trava a ingestão
                log.warning("Pink Sheet: data de atualização em formato inesperado: %r", achado.group(1))
            break

    cabecalho = next((i for i, linha in enumerate(linhas[:20]) if any(_nome_coluna(c) in config.PINK_SERIES for c in linha)), None)
    if cabecalho is None or cabecalho + 1 >= len(linhas):
        raise ErroFonte("não achei a linha com os nomes das séries na aba da Pink Sheet")
    nomes, unidades = linhas[cabecalho], linhas[cabecalho + 1]
    colunas = {}
    for j, nome in enumerate(nomes):
        serie = config.PINK_SERIES.get(_nome_coluna(nome))
        if serie is None:
            continue
        if serie in colunas:
            raise ErroFonte(f"a coluna {_nome_coluna(nome)!r} aparece duas vezes na Pink Sheet")
        unidade = str(unidades[j] or "").strip() if j < len(unidades) else ""
        if unidade != "($/mt)":
            raise ErroFonte(f"a unidade de {_nome_coluna(nome)!r} na Pink Sheet é {unidade!r}, e não ($/mt)")
        colunas[serie] = j
    faltando = [nome for nome, serie in config.PINK_SERIES.items() if serie not in colunas and serie in config.PINK_OBRIGATORIAS]
    if faltando:
        raise ErroFonte(f"a Pink Sheet não tem as colunas {faltando}; o layout pode ter mudado")
    for nome, serie in config.PINK_SERIES.items():
        if serie not in colunas:  # só ingerida, não comparada: o resto segue sem ela
            log.warning("Pink Sheet sem a coluna %r; ela fica vazia nesta ingestão.", nome)

    registros = []
    dados = linhas[cabecalho + 2 :]
    for n, linha in enumerate(dados):
        periodo = str(linha[0] or "").strip() if linha else ""
        if not periodo:
            continue
        achado = _RE_PERIODO_PINK.match(periodo)
        if not achado:
            # Nota no rodapé da aba, depois dos dados: para aí. Período estranho no meio é erro.
            depois = (str(r[0] or "").strip() for r in dados[n + 1 :] if r)
            if registros and not any(_RE_PERIODO_PINK.match(d) for d in depois):
                break
            raise ErroFonte(f"período inesperado na Pink Sheet: {periodo!r}")
        registro = {"mes": f"{achado.group(1)}-{achado.group(2)}"}
        for serie in config.PINK_SERIES.values():
            j = colunas.get(serie)
            registro[serie] = _numero(linha[j] if j is not None and j < len(linha) else None)
        registros.append(registro)
    df = pd.DataFrame(registros, columns=COLUNAS_PINK)
    if df.empty:
        raise ErroFonte("a aba da Pink Sheet não tem nenhum mês")
    if df["mes"].duplicated().any() or not df["mes"].is_monotonic_increasing:
        raise ErroFonte("os meses da Pink Sheet estão repetidos ou fora de ordem")
    return df, atualizado


def _numero(valor) -> float:
    """Número da célula; texto numérico ("170") também vale. "..", "…" e vazio viram NaN."""
    if isinstance(valor, bool):
        return float("nan")
    if isinstance(valor, (int, float)):
        return float(valor)
    try:
        return float(str(valor).strip().replace(",", "")) if valor is not None else float("nan")
    except ValueError:
        return float("nan")


def ultimo_mes_completo(df: pd.DataFrame) -> str | None:
    """Último mês em que as séries comparadas com o Brasil têm valor."""
    completos = df.dropna(subset=list(config.PINK_OBRIGATORIAS))
    return None if completos.empty else str(completos["mes"].iloc[-1])


def _fonte_gravada(raw_dir: Path) -> dict:
    try:
        return json.loads((Path(raw_dir) / config.PINK_FONTE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def ler_pink(raw_dir: Path = config.RAW_BENCHMARK_DIR) -> pd.DataFrame:
    caminho = Path(raw_dir) / config.PINK_ARQUIVO
    if not caminho.exists():
        return pd.DataFrame(columns=COLUNAS_PINK)
    return pd.read_csv(caminho, dtype={"mes": str})


def ultimo_pink_gravado(raw_dir: Path = config.RAW_BENCHMARK_DIR) -> str | None:
    df = ler_pink(raw_dir)
    return None if df.empty else ultimo_mes_completo(df)


def revisoes(antigo: pd.DataFrame, novo: pd.DataFrame) -> list[tuple[str, str, float, float]]:
    """(mês, série, valor gravado, valor novo) dos valores que o Banco Mundial mudou."""
    if antigo.empty:
        return []
    juntos = antigo.merge(novo, on="mes", suffixes=("_antigo", "_novo"))
    mudancas = []
    for serie in config.PINK_SERIES.values():
        a, n = juntos[f"{serie}_antigo"], juntos[f"{serie}_novo"]
        diferente = ~((a - n).abs() <= 1e-9) & ~(a.isna() & n.isna())
        for _, linha in juntos[diferente].iterrows():
            mudancas.append((linha["mes"], serie, linha[f"{serie}_antigo"], linha[f"{serie}_novo"]))
    return sorted(mudancas)


def ingerir_pink(
    cliente: ClienteFontes,
    inicio: str = config.INICIO_FASE6,
    fim: str | None = None,
    *,
    raw_dir: Path = config.RAW_BENCHMARK_DIR,
    forcar: bool = False,
) -> tuple[str, str, list]:
    """Baixa a Pink Sheet e regrava a série de `inicio` (ou do primeiro mês já gravado) até o
    último mês publicado (ou `fim`). Devolve (início, fim, revisões).

    Sem `forcar`, não baixa nada se `fim` já está gravado. Mês que falta no meio do período, ou
    um arquivo que termina antes do que já está gravado, é ErroFonte e nada é gravado.
    """
    gravado = ler_pink(raw_dir)
    ultimo_gravado = None if gravado.empty else ultimo_mes_completo(gravado)
    if not forcar and fim is not None and ultimo_gravado is not None and ultimo_gravado >= fim:
        log.info("Pink Sheet: nada a baixar, %s já está gravado em %s.", fim, raw_dir)
        return inicio, fim, []
    if fim is not None and ultimo_gravado is not None and fim < ultimo_gravado:
        # O arquivo vem sempre inteiro e é regravado inteiro: um --fim antes do gravado (para
        # rebaixar revisões do Comex Stat, por exemplo) não pode apagar os meses seguintes.
        log.info("Pink Sheet: --fim %s é anterior ao gravado; regravando até %s.", fim, ultimo_gravado)
        fim = ultimo_gravado
    html = cliente.get(config.PINK_PAGINA, esperado="html").text
    url = url_do_arquivo_pink(html)
    log.info("Pink Sheet: baixando %s", url)
    conteudo = cliente.get(url, esperado="xlsx").content
    df, atualizado = ler_xlsx_pink(conteudo)
    anterior = _fonte_gravada(raw_dir).get("atualizado_em")
    if atualizado and anterior and atualizado < anterior:
        raise ErroFonte(
            f"a edição baixada da Pink Sheet ({atualizado}) é anterior à gravada ({anterior}): um cache antigo"
            " desfaria as revisões. Nada foi gravado."
        )
    publicado = ultimo_mes_completo(df)
    if publicado is None:
        raise ErroFonte("a Pink Sheet não tem nenhum mês com ureia, DAP, TSP e KCl")
    if fim is None or fim > publicado:
        if fim is not None:
            log.info("Pink Sheet: %s ainda não foi publicado; indo até %s (edição de %s).", fim, publicado, atualizado or "data desconhecida")
        fim = publicado
    if ultimo_gravado is not None and ultimo_gravado > fim:
        raise ErroFonte(
            f"a Pink Sheet baixada vai até {fim}, antes do que já está gravado ({ultimo_gravado}). Nada foi"
            " gravado; confira o arquivo publicado."
        )
    if not gravado.empty:
        inicio = min(inicio, str(gravado["mes"].iloc[0]))
    periodo = df[(df["mes"] >= inicio) & (df["mes"] <= fim)].reset_index(drop=True)
    esperados = _meses(inicio, fim)
    faltam = sorted(set(esperados) - set(periodo["mes"]))
    lacunas = {s: periodo.loc[periodo[s].isna(), "mes"].tolist() for s in config.PINK_OBRIGATORIAS}
    lacunas = {s: m for s, m in lacunas.items() if m}
    if faltam or lacunas:
        detalhes = [f"sem a linha de {', '.join(faltam[:6])}"] if faltam else []
        detalhes += [f"{s} sem valor em {', '.join(m[:6])}" for s, m in lacunas.items()]
        raise ErroFonte(f"Pink Sheet com lacuna entre {inicio} e {fim}: {'; '.join(detalhes)}. Nada foi gravado.")
    mudancas = revisoes(gravado, periodo)
    for mes, serie, antes, depois in mudancas[:20]:
        log.warning("Pink Sheet revisou %s de %s: %s -> %s US$/t", serie, mes, antes, depois)
    if len(mudancas) > 20:
        log.warning("Pink Sheet: mais %d valores revisados.", len(mudancas) - 20)
    _gravar_csv(periodo[COLUNAS_PINK], Path(raw_dir) / config.PINK_ARQUIVO)
    fonte = {
        "fonte": "Banco Mundial, World Bank Commodity Price Data (The Pink Sheet)",
        "licenca": "CC BY 4.0",
        "pagina": config.PINK_PAGINA,
        "url": url,
        "atualizado_em": atualizado,
        "sha256": hashlib.sha256(conteudo).hexdigest(),
        "aba": config.PINK_ABA,
        "unidade": "US$ por tonelada métrica, nominal",
        "colunas": {serie: nome for nome, serie in config.PINK_SERIES.items()},
    }
    _gravar_texto(Path(raw_dir) / config.PINK_FONTE, json.dumps(fonte, ensure_ascii=False, indent=1) + "\n")
    log.info("Pink Sheet: %d meses gravados (%s a %s), %d valores revisados.", len(periodo), inicio, fim, len(mudancas))
    return inicio, fim, mudancas


def resumo_pink(df: pd.DataFrame, inicio: str, fim: str, mudancas: list | None = None) -> str:
    df = df[(df["mes"] >= inicio) & (df["mes"] <= fim)]
    if df.empty:
        return f"Pink Sheet: sem dados gravados para {inicio} a {fim}."
    ultimo = df.iloc[-1]
    valores = " | ".join(f"{s} {ultimo[s]:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") for s in config.PINK_SERIES.values() if pd.notna(ultimo[s]))
    linhas = [
        f"Resumo Pink Sheet (Banco Mundial, CC BY 4.0) {inicio} a {fim}",
        f"  Meses: {len(df)} | {ultimo['mes']} em US$/t: {valores}",
    ]
    if mudancas:
        linhas.append(f"  Valores revisados pelo Banco Mundial nesta ingestão: {len(mudancas)}")
    return "\n".join(linhas)


# --- BCB: PTAX (SGS 3698) ---------------------------------------------------------------


def _url_sgs(sufixo: str = "") -> str:
    return config.SGS_URL.format(codigo=config.SGS_PTAX_VENDA_MEDIA_MENSAL) + sufixo


def _mes_sgs(data: str) -> str:
    achado = _RE_DATA_SGS.match(str(data))
    if not achado:
        raise ErroFonte(f"data inesperada no SGS: {data!r} (esperado 01/MM/AAAA)")
    return f"{achado.group(2)}-{achado.group(1)}"


def normalizar_ptax(registros: list[dict]) -> pd.DataFrame:
    linhas = []
    for r in registros:
        try:
            valor = float(str(r["valor"]))
        except (KeyError, TypeError, ValueError) as exc:
            raise ErroFonte(f"valor inesperado no SGS: {r!r}") from exc
        if not 0 < valor < 100:
            raise ErroFonte(f"PTAX fora do intervalo plausível no SGS: {r!r}")
        linhas.append({"mes": _mes_sgs(r.get("data")), "brl_por_usd": valor})
    df = pd.DataFrame(linhas, columns=COLUNAS_PTAX)
    if df["mes"].duplicated().any():
        raise ErroFonte("o SGS devolveu meses repetidos")
    return df


def ultimo_ptax_publicado(cliente: ClienteFontes) -> str:
    registros = cliente.get(_url_sgs("/ultimos/1"), esperado="json", params={"formato": "json"}).json()
    if not registros:
        raise ErroFonte(f"o SGS não devolveu o último valor da série {config.SGS_PTAX_VENDA_MEDIA_MENSAL}")
    return _mes_sgs(registros[-1].get("data"))


def consultar_ptax(cliente: ClienteFontes, de: str, ate: str) -> pd.DataFrame:
    """Meses de `de` a `ate` (AAAA-MM): do dia 1 de `de` ao último dia de `ate`."""
    ano, mes = map(int, ate.split("-"))
    ultimo_dia = (pd.Timestamp(year=ano, month=mes, day=1) + pd.offsets.MonthEnd(0)).day
    params = {"formato": "json", "dataInicial": f"01/{de[5:7]}/{de[:4]}", "dataFinal": f"{ultimo_dia:02d}/{mes:02d}/{ano:04d}"}
    return normalizar_ptax(cliente.get(_url_sgs(), esperado="json", params=params).json())


def ler_ptax(raw_dir: Path = config.RAW_CAMBIO_DIR) -> pd.DataFrame:
    caminho = Path(raw_dir) / config.PTAX_ARQUIVO
    if not caminho.exists():
        return pd.DataFrame(columns=COLUNAS_PTAX)
    return pd.read_csv(caminho, dtype={"mes": str})


def ultimo_ptax_gravado(raw_dir: Path = config.RAW_CAMBIO_DIR) -> str | None:
    df = ler_ptax(raw_dir)
    return None if df.empty else str(df["mes"].max())


def ingerir_ptax(
    cliente: ClienteFontes,
    inicio: str = config.INICIO_FASE6,
    fim: str | None = None,
    *,
    raw_dir: Path = config.RAW_CAMBIO_DIR,
    forcar: bool = False,
    ultimo: str | None = None,
) -> tuple[str, str]:
    """Baixa os meses que faltam de [inicio, fim], em janelas de até cinco anos, e devolve o
    período efetivo. Mês já publicado que não vem na resposta é ErroFonte e nada é gravado."""
    gravado = ler_ptax(raw_dir)
    if not forcar and fim is not None and set(_meses(inicio, fim)) <= set(gravado["mes"]):
        log.info("PTAX: nada a baixar, %s a %s já está em %s.", inicio, fim, raw_dir)
        return inicio, fim
    if ultimo is None:
        ultimo = ultimo_ptax_publicado(cliente)
    if fim is None or fim > ultimo:
        if fim is not None:
            log.info("PTAX: %s ainda não foi publicado; indo até %s.", fim, ultimo)
        fim = ultimo
    if inicio > fim:
        log.info("PTAX: nada a fazer, %s é posterior ao último mês publicado (%s).", inicio, fim)
        return inicio, fim
    faltam = [m for m in _meses(inicio, fim) if forcar or m not in set(gravado["mes"])]
    if not faltam:
        log.info("PTAX: nada a baixar, %s a %s já está em %s.", inicio, fim, raw_dir)
        return inicio, fim
    novos = []
    for janela in _janelas(faltam, 12 * config.SGS_JANELA_ANOS):
        log.info("PTAX: consultando %s a %s...", janela[0], janela[-1])
        df = consultar_ptax(cliente, janela[0], janela[-1])
        novos.append(df[df["mes"].isin(janela)])
    novo = pd.concat(novos, ignore_index=True)
    sem_dado = sorted(set(faltam) - set(novo["mes"]))
    if sem_dado:
        raise ErroFonte(
            f"PTAX: o BCB já publicou até {ultimo}, mas não devolveu {', '.join(sem_dado[:12])}."
            " Nada foi gravado; tente de novo mais tarde."
        )
    mantidos = gravado[~gravado["mes"].isin(novo["mes"])]
    # Sem nada gravado ainda, só o novo: concat com DataFrame vazio é FutureWarning no pandas 2.
    junto = pd.concat([mantidos, novo], ignore_index=True) if not mantidos.empty else novo
    junto = junto.sort_values("mes", ignore_index=True)[COLUNAS_PTAX]
    _gravar_csv(junto, Path(raw_dir) / config.PTAX_ARQUIVO)
    log.info("PTAX: +%d meses (%s a %s).", len(novo), novo["mes"].min(), novo["mes"].max())
    return inicio, fim


def _janelas(meses: list[str], tamanho: int) -> list[list[str]]:
    """Trechos de meses seguidos, cada um com até `tamanho` meses: 2015 e 2026-09 faltando
    viram duas consultas, e não uma de onze anos."""
    trechos: list[list[str]] = []
    for mes in meses:
        if trechos and len(trechos[-1]) < tamanho and _meses(trechos[-1][-1], mes)[1:2] == [mes]:
            trechos[-1].append(mes)
        else:
            trechos.append([mes])
    return trechos


def resumo_ptax(df: pd.DataFrame, inicio: str, fim: str) -> str:
    df = df[(df["mes"] >= inicio) & (df["mes"] <= fim)]
    if df.empty:
        return f"PTAX: sem dados gravados para {inicio} a {fim}."
    ultimo = df.iloc[-1]
    return "\n".join(
        [
            f"Resumo PTAX venda, média mensal (BCB, SGS {config.SGS_PTAX_VENDA_MEDIA_MENSAL}) {inicio} a {fim}",
            f"  Meses: {len(df)} | {ultimo['mes']}: R$ {ultimo['brl_por_usd']:.4f} por US$".replace(".", ","),
        ]
    )
